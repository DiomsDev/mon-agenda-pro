from django.contrib.auth.decorators import login_required
from django.shortcuts import render, redirect, get_object_or_404
from django.utils import timezone

from .models import Activite, RendezVous
from missions.models import Mission
from .forms import ActiviteForm, RendezVousForm
from missions.forms import MissionForm

import calendar
from historique.models import HistoriqueAction
from historique.utils import enregistrer_action

from django.http import HttpResponse
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from docx import Document

from django.conf import settings


from django.db.models.functions import TruncMonth
from django.db.models import Count
import json
from datetime import timedelta


from django.http import JsonResponse


from historique.utils import creer_notification
from comptes.models import Utilisateur

from historique.models import Notification, Rappel
# =========================================================
# TABLEAU DE BORD
# =========================================================

@login_required
def dashboard(request):
    aujourd_hui = timezone.now().date()

    # Entreprise de l'utilisateur connecté
    entreprise = request.user.entreprise

    # Activités de l'entreprise
    activites_a_venir = Activite.objects.filter(
        entreprise=entreprise,
        date__gte=aujourd_hui
    )

    activites_effectuees = Activite.objects.filter(
        entreprise=entreprise,
        statut="effectuee"
    )

    # Rendez-vous de l'entreprise
    nombre_rendez_vous = RendezVous.objects.filter(
        entreprise=entreprise,
        statut__in=["planifie", "confirme"]
    ).count()

    # Missions de l'entreprise
    nombre_missions = Mission.objects.filter(
        entreprise=entreprise
    ).exclude(
        statut__in=["terminee", "annulee"]
    ).count()

    return render(
        request,
        "agenda/dashboard.html",
        {
            "activites_a_venir": activites_a_venir,
            "activites_effectuees": activites_effectuees,
            "nombre_rendez_vous": nombre_rendez_vous,
            "nombre_missions": nombre_missions,
        }
    )


# =========================================================
# AGENDA
# =========================================================

@login_required
def agenda(request):
    aujourdhui = timezone.localdate()

    # Entreprise de l'utilisateur connecté
    entreprise = request.user.entreprise

    annee = int(request.GET.get("annee", aujourdhui.year))
    mois = int(request.GET.get("mois", aujourdhui.month))

    cal = calendar.Calendar(firstweekday=0)
    semaines = cal.monthdatescalendar(annee, mois)

    # Seulement les activités de l'entreprise de l'utilisateur
    activites = Activite.objects.filter(
        entreprise=entreprise,
        date__year=annee,
        date__month=mois
    )

    # Seulement les rendez-vous de l'entreprise de l'utilisateur
    rendez_vous = RendezVous.objects.filter(
        entreprise=entreprise,
        date__year=annee,
        date__month=mois
    )

    evenements_par_jour = {}

    for a in activites:
        evenements_par_jour.setdefault(a.date, []).append({
            "type": "activite",
            "objet": a
        })

    for r in rendez_vous:
        evenements_par_jour.setdefault(r.date, []).append({
            "type": "rendez_vous",
            "objet": r
        })

    semaines_avec_evenements = []

    for semaine in semaines:
        jours = []

        for jour in semaine:
            jours.append({
                "date": jour,
                "dans_le_mois": jour.month == mois,
                "aujourdhui": jour == aujourdhui,
                "evenements": evenements_par_jour.get(jour, []),
            })

        semaines_avec_evenements.append(jours)

    mois_precedent = mois - 1 if mois > 1 else 12
    annee_mois_precedent = annee if mois > 1 else annee - 1

    mois_suivant = mois + 1 if mois < 12 else 1
    annee_mois_suivant = annee if mois < 12 else annee + 1

    context = {
        "semaines": semaines_avec_evenements,
        "nom_mois": calendar.month_name[mois].capitalize(),
        "annee": annee,
        "mois_precedent": mois_precedent,
        "annee_mois_precedent": annee_mois_precedent,
        "mois_suivant": mois_suivant,
        "annee_mois_suivant": annee_mois_suivant,
    }

    return render(request, "agenda/agenda.html", context)



@login_required
def activites(request):
    entreprise = request.user.entreprise

    activites = Activite.objects.filter(
          entreprise=entreprise
    ).order_by("-date", "-heure_debut")

    return render(
        request,
        "agenda/activites.html",
        {"activites": activites}
    )


@login_required
def missions(request):
    entreprise = request.user.entreprise

    missions = Mission.objects.filter(
        entreprise=entreprise
    ).order_by("-date_depart")

    return render(
        request,
        "agenda/missions.html",
        {"missions": missions}
    )
    

@login_required
def rendez_vous(request):
    entreprise = request.user.entreprise

    rendez_vous = RendezVous.objects.filter(
        entreprise=entreprise
    ).order_by("-date", "-heure")

    return render(
        request,
        "agenda/rendez_vous.html",
        {"rendez_vous": rendez_vous}
    )

@login_required
def visiteurs(request):
    return render(request, "agenda/visiteurs.html")


@login_required
def contacts(request):
    return render(request, "agenda/contacts.html")


@login_required
def taches(request):
    return render(request, "agenda/taches.html")


@login_required
def comptes_rendus(request):
    return render(request, "agenda/comptes_rendus.html")


@login_required
def documents(request):
    return render(request, "agenda/documents.html")

@login_required
def rappels(request):
    return render(request, "agenda/rappels.html")


@login_required
def notifications(request):
    mes_notifications = Notification.objects.filter(destinataire=request.user)
    mes_notifications.filter(lue=False).update(lue=True)
    return render(request, "agenda/notifications.html", {"notifications": mes_notifications})



@login_required
def reunions(request):
    return render(request, "agenda/reunions.html")

@login_required
def statistiques(request):
    aujourdhui = timezone.localdate()
    annee_courante = aujourdhui.year
    debut_semaine = aujourdhui - timedelta(days=aujourdhui.weekday())
    fin_semaine = debut_semaine + timedelta(days=6)
    debut_mois = aujourdhui.replace(day=1)

    # ===== Compteurs "en temps réel" =====
    stats = {
        "activites": {
            "semaine": Activite.objects.filter(date__gte=debut_semaine, date__lte=fin_semaine).count(),
            "mois": Activite.objects.filter(date__gte=debut_mois, date__year=annee_courante, date__month=aujourdhui.month).count(),
            "annee": Activite.objects.filter(date__year=annee_courante).count(),
        },
        "missions": {
            "semaine": Mission.objects.filter(date_depart__gte=debut_semaine, date_depart__lte=fin_semaine).count(),
            "mois": Mission.objects.filter(date_depart__year=annee_courante, date_depart__month=aujourdhui.month).count(),
            "annee": Mission.objects.filter(date_depart__year=annee_courante).count(),
        },
        "rendez_vous": {
            "semaine": RendezVous.objects.filter(date__gte=debut_semaine, date__lte=fin_semaine).count(),
            "mois": RendezVous.objects.filter(date__year=annee_courante, date__month=aujourdhui.month).count(),
            "annee": RendezVous.objects.filter(date__year=annee_courante).count(),
        },
    }

    # ===== Données mensuelles pour les graphiques (année en cours) =====
    noms_mois = ["Jan", "Fév", "Mar", "Avr", "Mai", "Jun", "Jul", "Aoû", "Sep", "Oct", "Nov", "Déc"]

    def repartition_mensuelle(queryset, champ_date):
        compteurs = [0] * 12
        donnees = (
            queryset.filter(**{f"{champ_date}__year": annee_courante})
            .annotate(mois=TruncMonth(champ_date))
            .values("mois")
            .annotate(total=Count("id"))
        )
        for ligne in donnees:
            compteurs[ligne["mois"].month - 1] = ligne["total"]
        return compteurs

    graphique = {
        "labels": noms_mois,
        "activites": repartition_mensuelle(Activite.objects.all(), "date"),
        "missions": repartition_mensuelle(Mission.objects.all(), "date_depart"),
        "rendez_vous": repartition_mensuelle(RendezVous.objects.all(), "date"),
    }

    context = {
        "stats": stats,
        "annee_courante": annee_courante,
        "graphique_json": json.dumps(graphique),
    }
    return render(request, "agenda/statistiques.html", context)


# =========================================================
# HISTORIQUE
# =========================================================

@login_required
def historique(request):
    entreprise = request.user.entreprise
    actions = HistoriqueAction.objects.filter(entreprise=entreprise).order_by("-date")

    recherche = request.GET.get("recherche", "").strip()
    if recherche:
        actions = actions.filter(
            utilisateur__username__icontains=recherche
        ) | actions.filter(
            action__icontains=recherche
        ) | actions.filter(
            description__icontains=recherche
        )

    type_action = request.GET.get("type_action", "").strip()
    if type_action:
        actions = actions.filter(action=type_action)

    actions = actions[:200]

    types_actions = (
        HistoriqueAction.objects
        .filter(entreprise=entreprise)
        .values_list("action", flat=True)
        .distinct()
        .order_by("action")
    )

    context = {
        "actions": actions,
        "recherche": recherche,
        "type_action": type_action,
        "types_actions": types_actions,
    }

    return render(request, "agenda/historique.html", context)

# =========================================================
# ASSISTANTS
# =========================================================

@login_required
def assistants(request):
    return render(request, "agenda/assistants.html")


# =========================================================
# PARAMÈTRES
# =========================================================

@login_required
def parametres(request):
    return render(request, "agenda/parametres.html")


# =========================================================
# ACTIVITÉS — CRUD
# =========================================================


@login_required
def nouvelle_activite(request):
    if request.method == "POST":
        form = ActiviteForm(request.POST)
        if form.is_valid():
            activite = form.save(commit=False)
            activite.cree_par = request.user
            activite.entreprise = request.user.entreprise
            activite.save()
            enregistrer_action(request.user, "Création d'activité", f"Activité « {activite.objet} » créée")

            if activite.statut == "en_attente":
                directeurs = Utilisateur.objects.filter(
                    entreprise=request.user.entreprise, role="directeur"
                )
                for directeur in directeurs:
                    creer_notification(
                        directeur,
                        f"Nouvelle activité à valider : « {activite.objet} »",
                        lien="/activites/",
                    )

            return redirect("activites")
    else:
        form = ActiviteForm()

    return render(request, "agenda/nouvelle_activite.html", {"form": form})



@login_required
def modifier_activite(request, pk):
    activite = get_object_or_404(Activite, pk=pk, entreprise=request.user.entreprise)
    ancien_statut = activite.statut

    if request.method == "POST":
        form = ActiviteForm(request.POST, instance=activite)
        if form.is_valid():
            activite = form.save()
            enregistrer_action(request.user, "Modification d'activité", f"Activité « {activite.objet} » modifiée")

            if ancien_statut == "en_attente" and activite.statut in ["confirmee", "annulee"]:
                verbe = "confirmée" if activite.statut == "confirmee" else "refusée"
                creer_notification(
                    activite.cree_par,
                    f"Votre activité « {activite.objet} » a été {verbe}",
                    lien="/activites/",
                )

            return redirect("activites")
    else:
        form = ActiviteForm(instance=activite)

    return render(request, "agenda/nouvelle_activite.html", {"form": form, "modification": True})




@login_required
def supprimer_activite(request, pk):
    activite = get_object_or_404(Activite, pk=pk, entreprise=request.user.entreprise)

    if request.method == "POST":
        objet = activite.objet
        activite.delete()
        enregistrer_action(request.user, "Suppression d'activité", f"Activité « {objet} » supprimée")
        return redirect("activites")

    return render(request, "agenda/supprimer_activite.html", {"activite": activite})

# =========================================================
# MISSIONS — CRUD
# =========================================================


@login_required
def nouvelle_mission(request):
    if request.method == "POST":
        form = MissionForm(request.POST)
        if form.is_valid():
            mission = form.save(commit=False)
            mission.cree_par = request.user
            mission.entreprise = request.user.entreprise
            mission.save()
            enregistrer_action(request.user, "Création de mission", f"Mission « {mission.motif} » créée")

            if request.user.role == "assistant":
                directeurs = Utilisateur.objects.filter(
                    entreprise=request.user.entreprise, role="directeur"
                )
                for directeur in directeurs:
                    creer_notification(
                        directeur,
                        f"Nouvelle mission ajoutée : « {mission.motif} »",
                        lien="/missions/",
                    )

            return redirect("missions")
    else:
        form = MissionForm()

    return render(request, "agenda/nouvelle_mission.html", {"form": form})


@login_required
def modifier_mission(request, pk):
    mission = get_object_or_404(Mission, pk=pk, entreprise=request.user.entreprise)

    if request.method == "POST":
        form = MissionForm(request.POST, instance=mission)
        if form.is_valid():
            form.save()
            enregistrer_action(request.user, "Modification de mission", f"Mission « {mission.motif} » modifiée")
            return redirect("missions")
    else:
        form = MissionForm(instance=mission)

    return render(request, "agenda/nouvelle_mission.html", {"form": form, "modification": True})


@login_required
def supprimer_mission(request, pk):
    mission = get_object_or_404(Mission, pk=pk, entreprise=request.user.entreprise)

    if request.method == "POST":
        motif = mission.motif
        mission.delete()
        enregistrer_action(request.user, "Suppression de mission", f"Mission « {motif} » supprimée")
        return redirect("missions")

    return render(request, "agenda/supprimer_mission.html", {"mission": mission})

# =========================================================
# RENDEZ-VOUS — CRUD
# =========================================================

@login_required
def nouveau_rendez_vous(request):
    if request.method == "POST":
        form = RendezVousForm(request.POST)
        if form.is_valid():
            rdv = form.save(commit=False)
            rdv.cree_par = request.user
            rdv.entreprise = request.user.entreprise
            rdv.save()
            enregistrer_action(request.user, "Création de rendez-vous", f"Rendez-vous « {rdv.objet} » créé")

            if request.user.role == "assistant":
                directeurs = Utilisateur.objects.filter(
                    entreprise=request.user.entreprise, role="directeur"
                )
                for directeur in directeurs:
                    creer_notification(
                        directeur,
                        f"Nouveau rendez-vous ajouté : « {rdv.objet} »",
                        lien="/rendez-vous/",
                    )

            return redirect("rendez_vous")
    else:
        form = RendezVousForm()

    return render(request, "agenda/nouveau_rendez_vous.html", {"form": form})


@login_required
def modifier_rendez_vous(request, pk):
    rdv = get_object_or_404(RendezVous, pk=pk, entreprise=request.user.entreprise)

    if request.method == "POST":
        form = RendezVousForm(request.POST, instance=rdv)
        if form.is_valid():
            form.save()
            enregistrer_action(request.user, "Modification de rendez-vous", f"Rendez-vous « {rdv.objet} » modifié")
            return redirect("rendez_vous")
    else:
        form = RendezVousForm(instance=rdv)

    return render(request, "agenda/nouveau_rendez_vous.html", {"form": form, "modification": True})


@login_required
def supprimer_rendez_vous(request, pk):
    rdv = get_object_or_404(RendezVous, pk=pk, entreprise=request.user.entreprise)

    if request.method == "POST":
        objet = rdv.objet
        rdv.delete()
        enregistrer_action(request.user, "Suppression de rendez-vous", f"Rendez-vous « {objet} » supprimé")
        return redirect("rendez_vous")

    return render(request, "agenda/supprimer_rendez_vous.html", {"rendez_vous": rdv})


# =========================================================
# COTE ARCHIVE
# =========================================================

@login_required
def archives(request):
    entreprise = request.user.entreprise

    activites_archivees = Activite.objects.filter(entreprise=entreprise, statut="effectuee").order_by("-date")
    missions_archivees = Mission.objects.filter(entreprise=entreprise, statut="terminee").order_by("-date_depart")
    rendez_vous_archives = RendezVous.objects.filter(entreprise=entreprise, statut__in=["termine", "annule"]).order_by("-date")

    context = {
        "activites_archivees": activites_archivees,
        "missions_archivees": missions_archivees,
        "rendez_vous_archives": rendez_vous_archives,
    }
    return render(request, "agenda/archives.html", context)




@login_required
def activite_pdf(request, pk):
    activite = get_object_or_404(Activite, pk=pk, entreprise=request.user.entreprise)

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="activite_{activite.pk}.pdf"'

    p = canvas.Canvas(response, pagesize=A4)
    largeur, hauteur = A4
    y = hauteur - 60

    p.setFont("Helvetica-Bold", 16)
    p.drawString(50, y, "DIRAGENDA — Fiche d'activité")
    y -= 40

    p.setFont("Helvetica-Bold", 12)
    p.drawString(50, y, "Objet :")
    p.setFont("Helvetica", 12)
    p.drawString(150, y, activite.objet)
    y -= 25

    champs = [
        ("Date", activite.date.strftime("%d/%m/%Y")),
        ("Date de fin", activite.date_fin.strftime("%d/%m/%Y") if activite.date_fin else "—"),
        ("Heure début", activite.heure_debut.strftime("%H:%M") if activite.heure_debut else "—"),
        ("Heure fin", activite.heure_fin.strftime("%H:%M") if activite.heure_fin else "—"),
        ("Lieu", activite.lieu or "—"),
        ("Contact", activite.contact or "—"),
        ("Statut", activite.get_statut_display()),
        ("Créé par", str(activite.cree_par)),
    ]

    p.setFont("Helvetica-Bold", 12)
    for label, valeur in champs:
        p.drawString(50, y, f"{label} :")
        p.setFont("Helvetica", 12)
        p.drawString(150, y, str(valeur))
        p.setFont("Helvetica-Bold", 12)
        y -= 22

    y -= 10
    p.drawString(50, y, "Observations :")
    p.setFont("Helvetica", 11)
    y -= 20
    for ligne in (activite.observations or "—").split("\n"):
        p.drawString(50, y, ligne)
        y -= 16

    p.showPage()
    p.save()
    return response


@login_required
def activite_word(request, pk):
    activite = get_object_or_404(Activite, pk=pk, entreprise=request.user.entreprise)

    document = Document()
    document.add_heading("DIRAGENDA — Fiche d'activité", level=1)

    table = document.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"

    lignes = [
        ("Objet", activite.objet),
        ("Date", activite.date.strftime("%d/%m/%Y")),
        ("Date de fin", activite.date_fin.strftime("%d/%m/%Y") if activite.date_fin else "—"),
        ("Heure début", activite.heure_debut.strftime("%H:%M") if activite.heure_debut else "—"),
        ("Heure fin", activite.heure_fin.strftime("%H:%M") if activite.heure_fin else "—"),
        ("Lieu", activite.lieu or "—"),
        ("Contact", activite.contact or "—"),
        ("Statut", activite.get_statut_display()),
        ("Créé par", str(activite.cree_par)),
    ]
    for label, valeur in lignes:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = str(valeur)

    document.add_heading("Observations", level=2)
    document.add_paragraph(activite.observations or "—")

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    response["Content-Disposition"] = f'attachment; filename="activite_{activite.pk}.docx"'
    document.save(response)
    return response


@login_required
def mission_pdf(request, pk):
    mission = get_object_or_404(Mission, pk=pk, entreprise=request.user.entreprise)

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="mission_{mission.pk}.pdf"'

    p = canvas.Canvas(response, pagesize=A4)
    largeur, hauteur = A4
    y = hauteur - 60

    p.setFont("Helvetica-Bold", 16)
    p.drawString(50, y, "DIRAGENDA — Fiche de mission")
    y -= 40

    champs = [
        ("Motif", mission.motif),
        ("Lieu", mission.lieu),
        ("Date de départ", mission.date_depart.strftime("%d/%m/%Y")),
        ("Date de retour", mission.date_retour.strftime("%d/%m/%Y")),
        ("Statut", mission.get_statut_display()),
        ("Créé par", str(mission.cree_par)),
    ]

    p.setFont("Helvetica-Bold", 12)
    for label, valeur in champs:
        p.drawString(50, y, f"{label} :")
        p.setFont("Helvetica", 12)
        p.drawString(180, y, str(valeur))
        p.setFont("Helvetica-Bold", 12)
        y -= 25

    p.showPage()
    p.save()
    return response


@login_required
def mission_word(request, pk):
    mission = get_object_or_404(Mission, pk=pk, entreprise=request.user.entreprise)

    document = Document()
    document.add_heading("DIRAGENDA — Fiche de mission", level=1)

    table = document.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"

    lignes = [
        ("Motif", mission.motif),
        ("Lieu", mission.lieu),
        ("Date de départ", mission.date_depart.strftime("%d/%m/%Y")),
        ("Date de retour", mission.date_retour.strftime("%d/%m/%Y")),
        ("Statut", mission.get_statut_display()),
        ("Créé par", str(mission.cree_par)),
    ]
    for label, valeur in lignes:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = str(valeur)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    response["Content-Disposition"] = f'attachment; filename="mission_{mission.pk}.docx"'
    document.save(response)
    return response



@login_required
def rendez_vous_pdf(request, pk):
    rdv = get_object_or_404(RendezVous, pk=pk, entreprise=request.user.entreprise)

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="rendez_vous_{rdv.pk}.pdf"'

    p = canvas.Canvas(response, pagesize=A4)
    largeur, hauteur = A4
    y = hauteur - 60

    p.setFont("Helvetica-Bold", 16)
    p.drawString(50, y, "DIRAGENDA — Fiche de rendez-vous")
    y -= 40

    champs = [
        ("Objet", rdv.objet),
        ("Date", rdv.date.strftime("%d/%m/%Y")),
        ("Heure", rdv.heure.strftime("%H:%M") if rdv.heure else "—"),
        ("Lieu", rdv.lieu or "—"),
        ("Personne", rdv.personne or "—"),
        ("Téléphone", rdv.telephone or "—"),
        ("Statut", rdv.get_statut_display()),
        ("Créé par", str(rdv.cree_par)),
    ]

    p.setFont("Helvetica-Bold", 12)
    for label, valeur in champs:
        p.drawString(50, y, f"{label} :")
        p.setFont("Helvetica", 12)
        p.drawString(180, y, str(valeur))
        p.setFont("Helvetica-Bold", 12)
        y -= 22

    y -= 10
    p.drawString(50, y, "Observations :")
    p.setFont("Helvetica", 11)
    y -= 20
    for ligne in (rdv.observations or "—").split("\n"):
        p.drawString(50, y, ligne)
        y -= 16

    p.showPage()
    p.save()
    return response


@login_required
def rendez_vous_word(request, pk):
    rdv = get_object_or_404(RendezVous, pk=pk, entreprise=request.user.entreprise)

    document = Document()
    document.add_heading("DIRAGENDA — Fiche de rendez-vous", level=1)

    table = document.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"

    lignes = [
        ("Objet", rdv.objet),
        ("Date", rdv.date.strftime("%d/%m/%Y")),
        ("Heure", rdv.heure.strftime("%H:%M") if rdv.heure else "—"),
        ("Lieu", rdv.lieu or "—"),
        ("Personne", rdv.personne or "—"),
        ("Téléphone", rdv.telephone or "—"),
        ("Statut", rdv.get_statut_display()),
        ("Créé par", str(rdv.cree_par)),
    ]
    for label, valeur in lignes:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = str(valeur)

    document.add_heading("Observations", level=2)
    document.add_paragraph(rdv.observations or "—")

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    response["Content-Disposition"] = f'attachment; filename="rendez_vous_{rdv.pk}.docx"'
    document.save(response)
    return response



def service_worker(request):
    with open(settings.BASE_DIR / "static" / "service-worker.js", "r") as f:
        contenu = f.read()
    return HttpResponse(contenu, content_type="application/javascript")




@login_required
def statistiques(request):
    entreprise = request.user.entreprise
    aujourdhui = timezone.localdate()
    annee_courante = aujourdhui.year
    debut_semaine = aujourdhui - timedelta(days=aujourdhui.weekday())
    fin_semaine = debut_semaine + timedelta(days=6)
    debut_mois = aujourdhui.replace(day=1)

    stats = {
        "activites": {
            "semaine": Activite.objects.filter(entreprise=entreprise, date__gte=debut_semaine, date__lte=fin_semaine).count(),
            "mois": Activite.objects.filter(entreprise=entreprise, date__gte=debut_mois, date__year=annee_courante, date__month=aujourdhui.month).count(),
            "annee": Activite.objects.filter(entreprise=entreprise, date__year=annee_courante).count(),
        },
        "missions": {
            "semaine": Mission.objects.filter(entreprise=entreprise, date_depart__gte=debut_semaine, date_depart__lte=fin_semaine).count(),
            "mois": Mission.objects.filter(entreprise=entreprise, date_depart__year=annee_courante, date_depart__month=aujourdhui.month).count(),
            "annee": Mission.objects.filter(entreprise=entreprise, date_depart__year=annee_courante).count(),
        },
        "rendez_vous": {
            "semaine": RendezVous.objects.filter(entreprise=entreprise, date__gte=debut_semaine, date__lte=fin_semaine).count(),
            "mois": RendezVous.objects.filter(entreprise=entreprise, date__year=annee_courante, date__month=aujourdhui.month).count(),
            "annee": RendezVous.objects.filter(entreprise=entreprise, date__year=annee_courante).count(),
        },
    }

    noms_mois = ["Jan", "Fév", "Mar", "Avr", "Mai", "Jun", "Jul", "Aoû", "Sep", "Oct", "Nov", "Déc"]

    def repartition_mensuelle(queryset, champ_date):
        compteurs = [0] * 12
        donnees = (
            queryset.filter(**{f"{champ_date}__year": annee_courante})
            .annotate(mois=TruncMonth(champ_date))
            .values("mois")
            .annotate(total=Count("id"))
        )
        for ligne in donnees:
            compteurs[ligne["mois"].month - 1] = ligne["total"]
        return compteurs

    graphique = {
        "labels": noms_mois,
        "activites": repartition_mensuelle(Activite.objects.filter(entreprise=entreprise), "date"),
        "missions": repartition_mensuelle(Mission.objects.filter(entreprise=entreprise), "date_depart"),
        "rendez_vous": repartition_mensuelle(RendezVous.objects.filter(entreprise=entreprise), "date"),
    }

    context = {
        "stats": stats,
        "annee_courante": annee_courante,
        "graphique_json": json.dumps(graphique),
    }
    return render(request, "agenda/statistiques.html", context)


@login_required
def statistiques_data(request):
    entreprise = request.user.entreprise
    aujourdhui = timezone.localdate()
    annee_courante = aujourdhui.year
    debut_semaine = aujourdhui - timedelta(days=aujourdhui.weekday())
    fin_semaine = debut_semaine + timedelta(days=6)
    debut_mois = aujourdhui.replace(day=1)

    stats = {
        "activites": {
            "semaine": Activite.objects.filter(entreprise=entreprise, date__gte=debut_semaine, date__lte=fin_semaine).count(),
            "mois": Activite.objects.filter(entreprise=entreprise, date__year=annee_courante, date__month=aujourdhui.month).count(),
            "annee": Activite.objects.filter(entreprise=entreprise, date__year=annee_courante).count(),
        },
        "missions": {
            "semaine": Mission.objects.filter(entreprise=entreprise, date_depart__gte=debut_semaine, date_depart__lte=fin_semaine).count(),
            "mois": Mission.objects.filter(entreprise=entreprise, date_depart__year=annee_courante, date_depart__month=aujourdhui.month).count(),
            "annee": Mission.objects.filter(entreprise=entreprise, date_depart__year=annee_courante).count(),
        },
        "rendez_vous": {
            "semaine": RendezVous.objects.filter(entreprise=entreprise, date__gte=debut_semaine, date__lte=fin_semaine).count(),
            "mois": RendezVous.objects.filter(entreprise=entreprise, date__year=annee_courante, date__month=aujourdhui.month).count(),
            "annee": RendezVous.objects.filter(entreprise=entreprise, date__year=annee_courante).count(),
        },
    }

    noms_mois = ["Jan", "Fév", "Mar", "Avr", "Mai", "Jun", "Jul", "Aoû", "Sep", "Oct", "Nov", "Déc"]

    def repartition_mensuelle(queryset, champ_date):
        compteurs = [0] * 12
        donnees = (
            queryset.filter(**{f"{champ_date}__year": annee_courante})
            .annotate(mois=TruncMonth(champ_date))
            .values("mois")
            .annotate(total=Count("id"))
        )
        for ligne in donnees:
            compteurs[ligne["mois"].month - 1] = ligne["total"]
        return compteurs

    graphique = {
        "labels": noms_mois,
        "activites": repartition_mensuelle(Activite.objects.filter(entreprise=entreprise), "date"),
        "missions": repartition_mensuelle(Mission.objects.filter(entreprise=entreprise), "date_depart"),
        "rendez_vous": repartition_mensuelle(RendezVous.objects.filter(entreprise=entreprise), "date"),
    }

    return JsonResponse({"stats": stats, "graphique": graphique})